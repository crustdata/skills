#!/usr/bin/env python3
"""Structural evals for the social-connections scripts. No network, no MCP.

    python3 evals/test_merge.py
    uv run --with openpyxl python evals/test_merge.py   # adds the workbook checks
"""
import importlib.util, json, subprocess, sys, tempfile
from collections import defaultdict
from pathlib import Path

SKILL = Path(__file__).resolve().parent.parent
MERGE = SKILL / "scripts" / "merge.py"
SHARD = SKILL / "scripts" / "shard_resolve.py"
SHEET = SKILL / "scripts" / "build_sheet.py"
HAVE_XLSX = importlib.util.find_spec("openpyxl") is not None

URN_A = "ACoAAAAAAAA1aaaaaaa"
URN_B = "ACoAAAAAAAA2bbbbbbb"
URN_C = "ACoAAAAAAAA3ccccccc"
URN_D = "ACoAAAAAAAA4ddddddd"
URN_E = "ACoAAAAAAAA5eeeeeee"
URN_F = "ACoAAAAAAAA6fffffff"


def url(slug):
    return f"https://www.linkedin.com/in/{slug}"


def rec(name, link, angle, strength="weak", evidence=("x",), **kw):
    return {"name": name, "linkedin_url": link, "angle": angle, "strength": strength,
            "evidence": list(evidence), **kw}


ANGLE_1 = [
    rec("Strong Both", url(URN_A), "team", "strong", ["same payments team 2016-2017"],
        current_company="Acme", current_title="TL"),
    rec("Weak Only", url(URN_C), "team", evidence=["same company, different function"],
        current_company="not listed on LinkedIn"),
    rec("Opted Out", url(URN_D), "team", evidence=["same company, different function"]),
]
ANGLE_2 = [
    # the same person as URN_A, seen by vanity url; only the resolver may join them
    rec("Strong Both", url("strong-both"), "posts", "strong", ["commented on their post"],
        urn=URN_A, current_company="Acme"),
    # a chatty angle: 6 weak items must not outscore a single strong item
    rec("Chatty", url(URN_B), "posts", evidence=[f"weak reason {i}" for i in range(6)],
        current_company="Google LLC"),
]
ANGLE_3 = [rec("Weak Only", url("weak-only"), "school", "medium", ["same lab 2014-2016"])]
DONE = [
    {"query_url": url(URN_A), "resolved_url": url("strong-both"), "current_company": "Acme",
     "current_title": "TL"},
    {"query_url": url(URN_C), "resolved_url": url("weak-only"),
     "current_company": "Stripe, Inc.", "current_title": "Engineer"},
    {"query_url": url(URN_D), "redacted": True},
]

FAILURES = []


def check(cond, msg):
    print(("  ok   " if cond else "  FAIL ") + msg)
    if not cond:
        FAILURES.append(msg)


def run(*args):
    return subprocess.run([sys.executable, *map(str, args)], capture_output=True, text=True)


def write(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data))


def rows_by_name(rd):
    by = defaultdict(list)
    for x in json.loads((rd / "final.json").read_text()):
        by[x["name"]].append(x)
    return by


def main_scenario():
    with tempfile.TemporaryDirectory() as td:
        rd = Path(td)
        write(rd / "run.json", {"name": "Test Target", "slug": "test-target"})
        write(rd / "raw" / "team.json", ANGLE_1)
        write(rd / "raw" / "posts.json", ANGLE_2)

        r = run(MERGE, rd)
        check(r.returncode == 0, "merge.py exits 0")
        by = rows_by_name(rd)
        check(len(by["Strong Both"]) == 2, "before resolution a urn-only row and a vanity row stay apart")
        check(by["Chatty"][0]["score"] <= 2, f"6 weak items from one angle capped at best 2 (score {by['Chatty'][0]['score']})")
        check(by["Chatty"][0]["tier"] in ("C", "D"), "chatty weak angle stays low tier")
        check(by["Chatty"][0]["employer_norm"] == "Google", "'Google LLC' normalises to 'Google'")
        check(by["Weak Only"][0]["employer_norm"] == "", "'not listed on LinkedIn' blanks the employer")
        check((rd / "final.md").read_text().startswith("# Test Target"), "final.md is titled from run.json")

        r2 = run(SHARD, rd, "--shards", "2")
        check(r2.returncode == 0, "shard_resolve.py exits 0")
        todo = [u for i in (0, 1) for u in json.loads((rd / "raw" / "resolve" / f"todo_{i}.json").read_text())]
        check(any(URN_C in u for u in todo), "a row with a placeholder employer is queued")
        check(any(URN_B in u for u in todo), "a urn-only row is queued even when it has an employer")

        write(rd / "raw" / "school.json", ANGLE_3)
        write(rd / "raw" / "resolve" / "done_0.json", DONE)
        check(run(MERGE, rd).returncode == 0, "re-merge exits 0")
        by = rows_by_name(rd)
        sb = by["Strong Both"]
        check(len(sb) == 1, f"the resolver joins the urn row to its vanity twin (got {len(sb)})")
        check(sb and sb[0]["tier"] == "A" and sb[0]["engaged"], "two strong items from two angles, one engaged -> tier A")
        check(sb and sb[0]["linkedin_url"] == url("strong-both"), "the joined row carries the vanity url")
        check(sb and sorted(sb[0]["angles"]) == ["posts", "team"], "angles merge across files")
        wo = by["Weak Only"]
        check(len(wo) == 1 and wo[0]["employer_norm"] == "Stripe" and sorted(wo[0]["angles"]) == ["school", "team"],
              "urn-only and vanity-only rows join via the resolver, with its employer")
        check("Opted Out" not in by, "a row the resolver found redacted is removed")

        check(run(SHEET, "--run-dir", rd, "--dry-run").returncode == 0, "build_sheet.py --dry-run needs no openpyxl")
        if HAVE_XLSX:
            from openpyxl import load_workbook
            n_d = sum(1 for rows in by.values() for x in rows if x["tier"] == "D")
            check(run(SHEET, "--run-dir", rd, "--no-tier-d").returncode == 0, "build_sheet.py writes the workbook")
            wb = load_workbook(rd / "test-target-connections.xlsx")
            check("tier D (appendix)" not in wb.sheetnames, "--no-tier-d drops the tier D tab")
            check(all(v[0] != "D" for v in wb["all people"].iter_rows(min_row=2, values_only=True)),
                  "--no-tier-d drops D rows from all people")
            check([v[1] for v in wb["summary"].iter_rows(values_only=True) if v and v[0] == "D"] == [n_d],
                  "the summary still counts tier D")
            check(wb["all people"].auto_filter.ref is not None, "people tabs carry a filter")


def stranger_scenario():
    """An agent pasted Carol's urn next to Bob's vanity url. Carol's own rows must not fold into Bob."""
    with tempfile.TemporaryDirectory() as td:
        rd = Path(td)
        write(rd / "raw" / "team.json", [rec("Bob Lee", url("bob-lee"), "team", urn=URN_E)])
        write(rd / "raw" / "posts.json", [rec("Carol Kim", url(URN_E), "posts", "strong", ["commented on the hiring post"])])
        write(rd / "raw" / "resolve" / "done_0.json", [{"query_url": url(URN_E), "resolved_url": url("carol-kim"),
                                                         "current_company": "Acme"}])
        check(run(MERGE, rd).returncode == 0, "stranger: merge.py exits 0")
        by = rows_by_name(rd)
        check(len(by["Bob Lee"]) == 1 and len(by["Carol Kim"]) == 1 and by["Bob Lee"][0]["tier"] != "A",
              "a copy-pasted urn does not hand one person's engagement to another")


def order_scenario():
    """File names and record order must not change the result."""
    people = [
        rec("Ann", url("ann"), "a", "strong", ["same team 2019"], current_company="Acme", current_title="PM"),
        rec("Ann", url("ann"), "b", "strong", ["same team 2019"], current_company="ACME", headline="h1"),
        rec("Ann", url("ann"), "c", "medium", ["same lab"], current_company="acme", headline="h2"),
        rec("Ben", url(URN_F), "a", "strong", ["reacted to 3 posts"], current_company="Plaid"),
        rec("Ben", url("ben"), "b", "weak", ["same school"], current_company="Stripe"),
        rec("Cy", url("cy"), "c", evidence=["x"], current_company="acme"),
    ]
    done = [{"query_url": url(URN_F), "error": "timeout"},
            {"query_url": url(URN_F), "resolved_url": url("ben"), "current_company": "Figma", "current_title": "Eng"}]
    outs = []
    for names, flip in ((("1.json", "2.json"), False), (("z.json", "a.json"), True)):
        with tempfile.TemporaryDirectory() as td:
            rd = Path(td)
            half = len(people) // 2
            parts = [people[:half], people[half:]]
            if flip:
                parts = [list(reversed(p)) for p in parts]
            write(rd / "raw" / names[0], parts[0])
            write(rd / "raw" / names[1], parts[1])
            write(rd / "raw" / "resolve" / ("done_0.json" if not flip else "done_0-retry.json"), done[1:])
            write(rd / "raw" / "resolve" / ("done_0-retry.json" if not flip else "done_0.json"), done[:1])
            run(MERGE, rd)
            outs.append((rd / "final.json").read_text())
    check(outs[0] == outs[1], "final.json does not depend on file names, record order or retry file order")
    ben = [x for x in json.loads(outs[0]) if x["name"] == "Ben"]
    check(ben and ben[0]["employer_norm"] == "Figma", "an answered resolver row beats an error row for the same url")


EDGE = [
    rec("John Smith", None, "panel", "strong", ["co-panelist at a 2019 summit"]),
    rec("John Smith", None, "patents", "strong", ["co-inventor on a 2020 patent"]),
    rec("Jurgen A", url("j%C3%BCrgen-m"), "team"),
    rec("Jurgen B", url("jürgen-m"), "school"),
    rec("Jurgen C", url("jürgen-m"), "posts"),
    rec("Query Url", url(URN_B) + "/?miniProfileUrn=abc", "posts", evidence=["reacted to a post"]),
    rec("Urn Only", None, "posts", evidence=["reacted to a post"], urn=URN_C),
    rec("Engines", url("engines"), "team", "strong", ["overlapped at Reaction Engines 2015-2017"]),
    rec("Tagged Co", url("tagged-co"), "team", "strong", ["overlapped at Tagged 2010-2012"]),
    rec("Headline Guess", url(URN_A), "posts", current_company="Google", current_title="SWE"),
    rec("Structured", url("structured"), "team", urn=URN_D, current_company="Stripe", current_title="Staff Engineer"),
    rec("Titled", url("titled"), "team", current_company="Acme", current_title="CTO"),
    rec("Moved", url("moved"), "team", current_company="Google", current_title="SWE"),
    rec("Lower", url("lower"), "team", current_company="stripe"),
    rec("Upper", url("upper"), "team", current_company="Stripe"),
    rec("Upper2", url("upper2"), "team", current_company="Stripe"),
    rec("Amzn", url("amzn"), "team", current_company="Amazon.com, Inc."),
    rec("Three", url("three"), "team", evidence=["team"]),
    rec("Three", url("three"), "school", evidence=["school"]),
    rec("Three", url("three"), "panel", evidence=["panel"]),
    rec("Twice", url("twice"), "team", "strong", ["same team 2019"]),
    rec("Twice", url("twice"), "team", "strong", ["same team 2019"]),
    rec("Engaged Strong", url("engaged-strong"), "posts", "strong", ["commented on their post"]),
    rec("Junk", url("junk"), "team", "STRONG", [None, 5]),
    rec("Numeric", url("numeric"), "team", 4, ["n"]),
    {"name": 5, "linkedin_url": url("int-name"), "current_company": {"x": 1}, "angle": 7, "evidence": ["i"]},
    rec("Surrogate", url("surrogate"), "team", evidence=["loved this \ud83d"], current_company="Acme \ud83d"),
    rec("Pub Url", "https://www.linkedin.com/pub/someone/1/2/3", "web"),
    rec('=HYPERLINK("http://x","y")', url("formula"), "team"),
    rec("Control", url("control"), "team", headline="tab\x0bhere", evidence=["esc \x1b here"]),
]
EDGE_DONE = [
    {"query_url": url(URN_A), "resolved_url": url("headline-guess"), "current_company": "Anthropic",
     "current_title": "Engineer"},
    {"query_url": url(URN_C), "redacted": True},
    {"query_url": url(URN_D), "resolved_url": url("structured"), "current_company": "Self employed",
     "current_title": ""},
    {"query_url": url("titled"), "current_company": "Acme Corporation", "current_title": "", "location": "Paris"},
    {"query_url": url("moved"), "current_company": "Stripe", "current_title": ""},
]


def edge_scenario():
    with tempfile.TemporaryDirectory() as td:
        rd = Path(td)
        write(rd / "run.json", {"name": "Tab\x0bTarget \ud83d", "slug": "edge"})
        write(rd / "raw" / "mixed.json", EDGE)
        write(rd / "raw" / "rowsenv.json", {"rows": [rec("Rows Env", url("rows-env"), None)]})
        write(rd / "raw" / "deepenv.json", {"ok": True, "value": {"results": [rec("Deep Env", url("deep-env"), None)]}})
        write(rd / "raw" / "single.json", rec("Single Rec", url("single-rec"), None))
        write(rd / "raw" / "nested.json", [[rec("Nested", url("nested"), "nested")]])
        write(rd / "raw" / "profile.json", {"name": "The Target", "jobs": [], "posts": []})
        write(rd / "raw" / "_scratch.json", [rec("Scratch", url("scratch"), "scratch")])
        # saved as a raw execute result rather than a bare list
        write(rd / "raw" / "resolve" / "done_0.json", {"ok": True, "value": EDGE_DONE})
        r = run(MERGE, rd)
        check(r.returncode == 0, f"edge: merge.py exits 0 on numbers, objects and surrogates {r.stderr[-200:]}")
        by = rows_by_name(rd)
        one = lambda n: by[n][0] if len(by[n]) == 1 else {}

        check(len(by["John Smith"]) == 2, "edge: unresolved namesakes in two angles stay two rows")
        jur = by["Jurgen A"] + by["Jurgen B"] + by["Jurgen C"]
        check(len(jur) == 1 and jur[0]["linkedin_url"] == url("j%C3%BCrgen-m"),
              "edge: percent-encoded, composed and decomposed slugs are one person with a quoted url")
        check(one("Query Url").get("linkedin_url") == url(URN_B), "edge: query string and trailing slash are stripped")
        check(not by["Urn Only"], "edge: a wrapped resolver file is read, so the redacted row goes")
        check(one("Engines").get("engaged") is False and one("Tagged Co").get("engaged") is False,
              "edge: an employer that is or contains an engagement verb is not engagement")
        check((one("Headline Guess").get("employer_norm"), one("Headline Guess").get("current_title")) == ("Anthropic", "Engineer"),
              "edge: the resolver's employer and title replace a headline guess")
        check((one("Structured").get("employer_norm"), one("Structured").get("current_title")) == ("Stripe", "Staff Engineer"),
              "edge: a 'Self employed' resolver answer does not overwrite a structured employer")
        t = one("Titled")
        check((t.get("current_title"), t.get("employer_norm"), t.get("location")) == ("CTO", "Acme", "Paris"),
              "edge: same employer with no resolver title keeps the title; the resolver fills location")
        check(one("Moved").get("current_title") == "", "edge: a new employer does not inherit the old employer's title")
        check({x["employer_norm"] for n in ("Lower", "Upper", "Upper2") for x in by[n]} == {"Stripe"},
              "edge: 'stripe' and 'Stripe' share one bucket, in the more common spelling")
        check(one("Amzn").get("employer_norm") == "Amazon", "edge: 'Amazon.com, Inc.' normalises to Amazon")
        check(one("Three").get("score") == 3, "edge: the best-two cap is per angle, not per person")
        check((one("Twice").get("score"), one("Twice").get("tier")) == (4, "B"), "edge: identical evidence counts once")
        check(one("Engaged Strong").get("tier") == "A" and one("Query Url").get("tier") == "B",
              "edge: engaged + strong is A, engagement alone is B")
        check(one("Junk").get("score") == 1, "edge: a row with no named reason scores as one weak item")
        check(one("5").get("angles") == ["7"], "edge: numeric name and angle become text")
        check(one("Pub Url").get("linkedin_url", "").startswith("https://www.linkedin.com/pub/"),
              "edge: a url with no canonical form is kept for the user")
        check(len(by["Rows Env"]) + len(by["Deep Env"]) + len(by["Single Rec"]) == 3,
              "edge: rows, nested value and single-record envelopes are read")
        check(one("Rows Env").get("angles") == ["rowsenv"], "edge: a record with no angle takes its file's name")
        check(not by["Scratch"] and not by["The Target"], "edge: _scratch files and profile dumps are not rows")
        check(len(by["Nested"]) == 1, "edge: a list of lists is flattened")

        check(run(SHARD, rd, "--shards", "1").returncode == 0, "edge: shard_resolve.py exits 0")
        todo = json.loads((rd / "raw" / "resolve" / "todo_0.json").read_text())
        check(todo and all(u.startswith("https://www.linkedin.com/in/") and "?" not in u for u in todo),
              "edge: only canonical profile urls are queued")

        if HAVE_XLSX:
            from openpyxl import load_workbook
            r = run(SHEET, "--run-dir", rd)
            check(r.returncode == 0, f"edge: control chars, surrogates and numeric fields still write a workbook {r.stderr[-200:]}")
            wb = load_workbook(rd / "edge-connections.xlsx")
            cells = [c for row in wb["all people"].iter_rows(min_row=2) for c in row]
            check(all(c.data_type != "f" for c in cells), "edge: formula-looking text stays text")


def shard_scenario():
    with tempfile.TemporaryDirectory() as td:
        rd = Path(td)
        write(rd / "final.json", [{"linkedin_url": "https://www.linkedin.com/pub/x/1/2/3", "employer_norm": ""},
                                  {"linkedin_url": "", "urn": "urn:li:member:1", "employer_norm": ""}])
        run(SHARD, rd)
        check(not list((rd / "raw" / "resolve").glob("todo_*.json")), "shard: nothing non-canonical is queued")


def main():
    main_scenario()
    stranger_scenario()
    order_scenario()
    edge_scenario()
    shard_scenario()
    if not HAVE_XLSX:
        print("  skip openpyxl not installed, workbook checks not run")
    print()
    if FAILURES:
        print(f"{len(FAILURES)} FAILED")
        sys.exit(1)
    print("all evals pass")


if __name__ == "__main__":
    main()
